"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

# Các cột công thức của mẫu — TUYỆT ĐỐI không ghi đè.
FORMULA_COLS = ("step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise")
MAX_DATA_ROW = 61          # sheet Experiments có sẵn 60 dòng dữ liệu (dòng 2..61)

_LOSS_NAME = {"ce": "CE", "mse": "MSE"}
_OPT_NAME = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _json_default(o):
    """Chuyển kiểu numpy/tuple về kiểu JSON gốc (json.dump không tự làm việc này)."""
    import numpy as np
    if isinstance(o, np.integer):
        return int(o)
    if isinstance(o, np.floating):
        return float(o)
    if isinstance(o, np.bool_):
        return bool(o)
    if isinstance(o, np.ndarray):
        return o.tolist()
    if isinstance(o, (tuple, set)):
        return list(o)
    raise TypeError(f"không serialize được: {type(o)}")


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file."""
    d = Path(results_dir)
    d.mkdir(parents=True, exist_ok=True)
    payload = {"cfg": result["cfg"], "history": result["history"], "summary": result["summary"]}
    path = d / f"{result['cfg']['exp_id']}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2, default=_json_default)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    d = Path(results_dir)
    out = []
    for p in sorted(d.glob("*.json")):
        with open(p, encoding="utf-8") as f:
            r = json.load(f)
        r.setdefault("cfg", {}).setdefault("exp_id", p.stem)
        out.append(r)
    out.sort(key=lambda r: r["cfg"]["exp_id"])
    return out


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary + eval (nếu có) + figure_file."""
    cfg, s = result["cfg"], result["summary"]
    clip = cfg.get("clip_norm")
    row = {
        "exp_id": cfg["exp_id"],
        "group": cfg.get("group", "other"),
        "description": cfg.get("description", ""),
        "loss": _LOSS_NAME.get(cfg["loss"], cfg["loss"]),
        "optimizer": _OPT_NAME.get(cfg["optimizer"], cfg["optimizer"]),
        "lr": cfg["lr"],
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": "-".join(str(h) for h in cfg["hidden"]),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": "none" if clip is None else float(clip),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg["seed"],
        "step0_loss": round(s["step0_loss"], 6),
        "best_val_loss": round(s["best_val_loss"], 6),
        "best_epoch": s["best_epoch"],
        "final_train_loss": round(s["final_train_loss"], 6),
        "final_val_loss": round(s["final_val_loss"], 6),
        "val_acc": round(s["val_acc"], 6),
        "val_macro_f1": round(s["val_macro_f1"], 6),
        "time_per_epoch_s": round(s["time_per_epoch_s"], 2),
        "peak_mem_MB": round(s["peak_mem_MB"], 1),
        "diverged": "Y" if s.get("diverged") else "N",
        "figure_file": f"figures/{cfg['exp_id']}.png",
        "notes": notes or s.get("notes", "") or "",
    }
    # eval_acc / eval_macro_f1: CHỈ điền cho baseline và cấu hình cuối cùng
    if eval_scores:
        row["eval_acc"] = round(float(eval_scores["accuracy"]), 6)
        row["eval_macro_f1"] = round(float(eval_scores["macro_f1"]), 6)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None,
               summary_notes: dict[str, str] | None = None) -> str:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Bỏ qua mọi cột công thức (FORMULA_COLS và bất kỳ ô nào đang chứa công thức "=").
    """
    import openpyxl

    if len(rows) > MAX_DATA_ROW - 1:
        raise ValueError(f"quá nhiều dòng: {len(rows)} > {MAX_DATA_ROW - 1}")

    wb = openpyxl.load_workbook(template_path)      # KHÔNG data_only=True (sẽ mất công thức)
    ws = wb["Experiments"]

    # map tên cột -> chỉ số cột (1-based)
    header = {c.value: c.column for c in ws[1] if c.value is not None}

    for i, row in enumerate(rows):
        r = i + 2
        for key, value in row.items():
            col = header.get(key)
            if col is None or key in FORMULA_COLS:
                continue
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                continue                              # ô công thức: không đụng
            cell.value = value

    # ---- sheet Seeds: các exp_id của lần chạy baseline (để tính mean / std / 2σ)
    if seed_ids:
        ws_s = wb["Seeds"]
        for i, eid in enumerate(seed_ids[:5]):
            ws_s.cell(row=2 + i, column=1).value = eid

    # ---- sheet Summary: nhận xét ngắn cho từng nhóm
    if summary_notes:
        ws_m = wb["Summary"]
        for i, row in enumerate(ws_m.iter_rows(min_row=2, max_col=1), start=2):
            g = row[0].value
            if g in summary_notes:
                ws_m.cell(row=i, column=8).value = summary_notes[g]

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"đã ghi {out_path} ({len(rows)} dòng thí nghiệm)")
    print("Mở file bằng Excel/LibreOffice để các cột công thức (2σ, beyond_noise) tính lại.")
    return out_path


def noise_stats(rows: list[dict], seed_ids: list[str]) -> dict:
    """Tính trung bình / độ lệch chuẩn mẫu / ngưỡng 2σ từ các dòng baseline (theo seed_ids).

    Dùng để quyết định một chênh lệch có phải bằng chứng hay chỉ là nhiễu (xem RUBRIC mục 4).
    """
    import numpy as np
    by_id = {r["exp_id"]: r for r in rows}
    picked = [by_id[e] for e in seed_ids if e in by_id]
    if not picked:
        raise ValueError("không tìm thấy dòng baseline nào trong seed_ids")
    out = {"n_seeds": len(picked), "seed_ids": [r["exp_id"] for r in picked]}
    for metric in ("val_acc", "val_macro_f1", "best_val_loss"):
        v = np.array([r[metric] for r in picked], dtype=float)
        out[metric] = dict(values=[float(x) for x in v], mean=float(v.mean()),
                           std=float(v.std(ddof=1)) if len(v) > 1 else float("nan"),
                           two_sigma=float(2 * v.std(ddof=1)) if len(v) > 1 else float("nan"))
    return out